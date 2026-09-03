#!/usr/bin/env python3
"""Regression tests for the Publication Structural Fidelity Gate."""

from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("publication_validator", ROOT / "scripts" / "validate_publication_structure.py")
validator = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
sys.modules[SPEC.name] = validator
SPEC.loader.exec_module(validator)


FIXTURE = """# Architecture Guide

## Concepts
### Participants
#### Physical Node
##### Identity

The node MUST expose a stable identity. Consumers SHOULD inspect freshness and MAY cache descriptions.

1. Discover the node
2. Read the manifest

1. Start a fresh independent list
2. Verify the result

- semantic contract
  - identity
  - authority
- runtime evidence

```python
def run():
    return "ok"
```

```text
Root
├── Node A -> State
└── Node B -> Action
```

> A quoted rationale.

| Field | Meaning |
| --- | --- |
| id | Stable identifier |
"""


def _add_num_instance(document: Document, num_id: int, kind: str = "ordered") -> None:
    numbering = document.part.numbering_part.element
    fmt = "decimal" if kind == "ordered" else "bullet"
    abstract = None
    for candidate in numbering.findall(qn("w:abstractNum")):
        level = candidate.find(qn("w:lvl"))
        num_fmt = level.find(qn("w:numFmt")) if level is not None else None
        if num_fmt is not None and num_fmt.get(qn("w:val")) == fmt:
            abstract = candidate.get(qn("w:abstractNumId"))
            break
    if abstract is None:
        raise AssertionError(f"no {fmt} abstract numbering definition")
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_id = OxmlElement("w:abstractNumId")
    abstract_id.set(qn("w:val"), abstract)
    num.append(abstract_id)
    numbering.append(num)


def _set_num(paragraph, num_id: int, level: int = 0) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    ppr.append(num_pr)


def _make_doc(path: Path, *, broken: str | None = None) -> None:
    document = Document()
    for style_name in ("Code", "ASCII Diagram"):
        style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "DejaVu Sans Mono"
    source = validator.parse_source(FIXTURE)
    for heading in source.headings:
        document.add_paragraph(heading.text, style=f"Heading {heading.level}")

    toc = document.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-3" \\h')
    toc._p.append(field)
    for heading in source.headings:
        if heading.level <= 3:
            document.add_paragraph(heading.text)
    document.add_page_break()

    document.add_paragraph("The node MUST expose a stable identity. Consumers SHOULD inspect freshness and MAY cache descriptions.")
    _add_num_instance(document, 101)
    _add_num_instance(document, 102)
    _add_num_instance(document, 103, "bullet")
    lists = [
        [("ordered", 0, "Discover the node"), ("ordered", 0, "Read the manifest")],
        [("ordered", 0, "Start a fresh independent list"), ("ordered", 0, "Verify the result")],
        [("unordered", 0, "semantic contract"), ("unordered", 1, "identity"), ("unordered", 1, "authority"), ("unordered", 0, "runtime evidence")],
    ]
    ids = [101, 101 if broken == "list-continuation" else 102, 103]
    for block, num_id in zip(lists, ids):
        for kind, level, text in block:
            paragraph = document.add_paragraph(text, style="List Number" if kind == "ordered" else "List Bullet")
            _set_num(paragraph, num_id, level)
        document.add_paragraph()

    code = document.add_paragraph(style="Code")
    code.add_run('def run():\n    return "ok"')
    diagram = document.add_paragraph(style="ASCII Diagram")
    diagram.add_run("Root\n├── Node A -> State\n└── Node B -> Action")
    document.add_paragraph("A quoted rationale.", style="Quote")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Meaning"
    table.cell(1, 0).text = "id"
    table.cell(1, 1).text = "Stable identifier"

    if broken == "h4-leak":
        document.add_paragraph("#### Physical Node")
    elif broken == "h5-leak":
        document.add_paragraph("##### Identity")
    elif broken == "empty-toc":
        for paragraph in document.paragraphs:
            if paragraph.text in [heading.text for heading in source.headings if heading.level <= 3]:
                paragraph.text = ""
    elif broken == "heading-order":
        document.paragraphs[3].style = "Heading 5"
    elif broken == "missing-table-row":
        table.rows[1].cells[1].text = ""
    document.save(path)


class StructuralValidatorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "fixture.md"
        self.source.write_text(FIXTURE, encoding="utf-8")
        self.profile = self.root / "profile.yaml"
        self.profile.write_text("toc:\n  enabled: true\n  depth: 3\n", encoding="utf-8")

    def tearDown(self):
        self.temp.cleanup()

    def result(self, broken=None):
        docx = self.root / f"{broken or 'valid'}.docx"
        _make_doc(docx, broken=broken)
        return validator.validate(self.source, docx, self.profile)

    def assert_invalid(self, broken, code):
        result = self.result(broken)
        self.assertEqual(result["status"], "INVALID")
        self.assertIn(code, {finding["code"] for finding in result["findings"]})

    def test_h1_h5_mapping_passes(self):
        result = self.result()
        self.assertEqual(result["status"], "VALID")
        self.assertEqual(result["source"]["H5"], result["docx"]["H5"])

    def test_literal_h4_leak_invalid(self):
        self.assert_invalid("h4-leak", "LITERAL_MARKDOWN_HEADING_LEAK")

    def test_literal_h5_leak_invalid(self):
        self.assert_invalid("h5-leak", "LITERAL_MARKDOWN_HEADING_LEAK")

    def test_separate_ordered_lists_restart(self):
        result = self.result()
        self.assertNotIn("UNEXPECTED_LIST_CONTINUATION", {f["code"] for f in result["findings"]})

    def test_unintended_continuation_invalid(self):
        self.assert_invalid("list-continuation", "UNEXPECTED_LIST_CONTINUATION")

    def test_nested_list_levels_pass(self):
        result = self.result()
        self.assertNotIn("LIST_LEVEL_MISMATCH", {f["code"] for f in result["findings"]})

    def test_code_indentation_passes(self):
        self.assertEqual(self.result()["status"], "VALID")

    def test_ascii_spacing_passes(self):
        self.assertEqual(self.result()["docx"]["ascii_diagram_style_blocks"], 1)

    def test_empty_toc_invalid(self):
        self.assert_invalid("empty-toc", "TOC_EMPTY")

    def test_useful_toc_passes(self):
        self.assertEqual(self.result()["docx"]["toc_entries"], 3)

    def test_heading_order_mismatch_invalid(self):
        self.assert_invalid("heading-order", "HEADING_LEVEL_MISMATCH")

    def test_heading_text_missing_invalid(self):
        path = self.root / "missing-heading.docx"
        _make_doc(path)
        document = Document(path)
        document.paragraphs[0].text = "Other title"
        document.save(path)
        result = validator.validate(self.source, path, self.profile)
        self.assertIn("HEADING_TEXT_MISSING", {f["code"] for f in result["findings"]})

    def test_source_hash_unchanged_passes(self):
        result = self.result()
        self.assertEqual(result["source_sha256"], validator.sha256(self.source))

    def test_source_hash_change_invalid(self):
        docx = self.root / "hash.docx"
        _make_doc(docx)
        self.assertEqual(validator.validate(self.source, docx, self.profile, "0" * 64)["status"], "INVALID")

    def test_table_row_loss_invalid(self):
        self.assert_invalid("missing-table-row", "TABLE_ROW_LOSS")

    def test_valid_long_technical_document_passes(self):
        self.assertEqual(self.result()["status"], "VALID")

    def test_generic_architecture_fixture_passes(self):
        result = self.result()
        self.assertEqual(result["source"]["fenced_code_blocks"], 2)
        self.assertEqual(result["status"], "VALID")

    def test_intentionally_broken_publication_invalid(self):
        self.assert_invalid("h4-leak", "LITERAL_MARKDOWN_HEADING_LEAK")

    def test_normative_keywords_preserved(self):
        self.assertEqual(self.result()["status"], "VALID")

    def test_code_and_diagram_are_distinct(self):
        result = self.result()
        self.assertEqual(result["docx"]["code_style_blocks"], 1)
        self.assertEqual(result["docx"]["ascii_diagram_style_blocks"], 1)

    def test_old_failure_cannot_be_candidate_ready(self):
        result = self.result("h4-leak")
        self.assertNotEqual(result["status"], "VALID")


if __name__ == "__main__":
    unittest.main()
