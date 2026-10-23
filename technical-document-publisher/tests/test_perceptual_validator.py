#!/usr/bin/env python3
"""Tests for the perceptual, black-and-white, and role-collapse signals."""

from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("perceptual_validator", ROOT / "scripts" / "validate_perceptual_publication.py")
perceptual = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
sys.modules[SPEC.name] = perceptual
SPEC.loader.exec_module(perceptual)


def _style(document: Document, name: str, size: float, *, bold: bool = False, before: float = 0, after: float = 4, indent: float = 0, border: bool = False, shading: bool = False, color: str | None = None):
    try:
        style = document.styles[name]
    except KeyError:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "DejaVu Sans"
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.left_indent = Cm(indent)
    ppr = style._element.get_or_add_pPr()
    if border:
        p_border = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "8")
        p_border.append(left)
        ppr.append(p_border)
    if shading:
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "F3F5F7")
        ppr.append(shade)
    return style


def _make_doc(path: Path, weak: str | None = None) -> None:
    d = Document()
    _style(d, "Normal", 10.5, after=5)
    _style(d, "Heading 2", 17, bold=True, before=16, after=8)
    _style(d, "Heading 3", 14, bold=True, before=12, after=6)
    topic_size = 10.5 if weak == "color-only" else (11 if weak == "topic" else 12)
    topic_border = weak not in {"topic", "color-only"}
    _style(d, "Heading 4", topic_size, bold=weak != "color-only", before=9 if topic_border else 0, after=5, indent=0.15 if topic_border else 0, border=topic_border, color="0000FF" if weak == "color-only" else None)
    label_border = weak != "label"
    _style(d, "Lead-in", 10.5, bold=True, before=4 if label_border else 0, after=5, indent=0.1 if label_border else 0, border=label_border, shading=label_border)
    for name in ("Runtime Label", "Question Label"):
        _style(d, name, 10.5, bold=True, after=5)
    d.add_paragraph("Chapter", style="Heading 2")
    d.add_paragraph("Major section", style="Heading 3")
    d.add_paragraph("Technical topic", style="Heading 4")
    d.add_paragraph("Local label", style="Lead-in")
    d.add_paragraph("Body text", style="Normal")
    d.save(path)


class PerceptualValidatorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.profile = self.root / "profile.yaml"
        self.profile.write_text("visual:\n  minimum_topic_body_distance: 0.45\n  minimum_local_label_body_distance: 0.38\n", encoding="utf-8")

    def tearDown(self):
        self.temp.cleanup()

    def report(self, weak=None):
        path = self.root / f"{weak or 'strong'}.docx"
        _make_doc(path, weak=weak)
        return perceptual.validate(path, self.profile)

    def test_strong_roles_pass_with_single_renderer_warning(self):
        result = self.report()
        self.assertEqual(result["status"], "VALID_WITH_WARNINGS")
        self.assertEqual(result["cross_renderer"], "CROSS_RENDERER_NOT_VERIFIED")

    def test_weak_topic_fails_perceptual_gate(self):
        codes = {item["code"] for item in self.report("topic")["findings"]}
        self.assertIn("TOPIC_BODY_CONTRAST_TOO_LOW", codes)
        self.assertIn("PERCEPTUAL_HIERARCHY_AMBIGUOUS", codes)

    def test_weak_local_label_fails_perceptual_gate(self):
        self.assertIn("LOCAL_LABEL_BODY_CONTRAST_TOO_LOW", {item["code"] for item in self.report("label")["findings"]})

    def test_color_only_topic_fails_black_white_signal(self):
        codes = {item["code"] for item in self.report("color-only")["findings"]}
        self.assertIn("TOPIC_BODY_CONTRAST_TOO_LOW", codes)

    def test_profile_can_map_chapters_to_heading_one(self):
        path = self.root / "report.docx"
        _make_doc(path)
        d = Document(path)
        _style(d, "Heading 1", 21, bold=True, before=20, after=10)
        d.save(path)
        self.profile.write_text(
            "roles:\n  chapter_style: Heading 1\n  major_section_style: Heading 2\n"
            "  technical_topic_style: Heading 3\n", encoding="utf-8")
        result = perceptual.validate(path, self.profile)
        self.assertEqual(result["roles"]["Chapter"]["name"], "Heading 1")
        self.assertEqual(result["roles"]["Technical Topic"]["name"], "Heading 3")

    def test_missing_selected_role_style_blocks_validation(self):
        path = self.root / "report.docx"
        _make_doc(path)
        self.profile.write_text("roles:\n  chapter_style: Missing Chapter\n", encoding="utf-8")
        result = perceptual.validate(path, self.profile)
        self.assertEqual(result["status"], "INVALID")
        self.assertIn("ROLE_STYLE_MISSING", {item["code"] for item in result["findings"]})


if __name__ == "__main__":
    unittest.main()
