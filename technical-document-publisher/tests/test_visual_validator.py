#!/usr/bin/env python3
"""Regression tests for the technical-whitepaper visual design gate."""

from __future__ import annotations

import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("visual_validator", ROOT / "scripts" / "validate_publication_visual_structure.py")
visual = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
sys.modules[SPEC.name] = visual
SPEC.loader.exec_module(visual)


def _style(document: Document, name: str, size: float, *, font: str = "DejaVu Sans", bold: bool | None = None, before: float = 0, after: float = 4, indent: float = 0):
    try:
        style = document.styles[name]
    except KeyError:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.left_indent = Cm(indent)
    return style


def _make_doc(path: Path, *, broken: str | None = None) -> None:
    document = Document()
    _style(document, "Normal", 10.5, after=5)
    heading_sizes = {1: 19, 2: 15.5, 3: 12.5, 4: 11, 5: 10.2, 6: 9.8}
    for level, size in heading_sizes.items():
        if broken == "flat" and level in (1, 2, 3):
            size = 10.5
        if broken == "inverted" and level == 1:
            size = 13
        heading_before = 16 if level == 1 else 10 if level == 2 else 7
        if broken == "spacing":
            heading_before = 0
        _style(document, f"Heading {level}", size, bold=True, before=heading_before, after=5)
    _style(document, "Code", 8.8, font="DejaVu Sans Mono", before=4, after=6)
    _style(document, "ASCII Diagram", 8.8 if broken != "ascii-body" else 10.5, font="DejaVu Sans Mono" if broken != "ascii-body" else "DejaVu Sans", before=4, after=6)
    for name in ("Lead-in", "Runtime Label", "Question Label"):
        _style(document, name, 10.5, bold=True, after=4)
    _style(document, "Table Header", 9.2, bold=True, after=2)
    _style(document, "Table Body", 9.2, after=2)
    for name in ("Bullet 1", "Bullet 2", "Numbered 1", "Numbered 2"):
        _style(document, name, 10.5, after=1)
    for level, indent in ((1, 0.2), (2, 0.6), (3, 1.0)):
        size = 9.5 - (level - 1) * 0.4 if broken != "flat-toc" else 9.5
        _style(document, f"TOC {level}", size, after=1, indent=indent if broken != "flat-toc" else 0.2)
    _style(document, "Header", 8, after=0)
    _style(document, "Footer", 8, after=0)
    document.add_paragraph("Architecture", style="Heading 1")
    document.add_paragraph("Concepts", style="Heading 2")
    document.add_paragraph("Node", style="Heading 3")
    document.add_paragraph("Body text.", style="Normal")
    document.add_paragraph("def observe():\n    return True", style="Code")
    document.add_paragraph("Root\n├── Node\n└── State", style="ASCII Diagram")
    document.add_paragraph("TOC 1", style="TOC 1")
    document.add_paragraph("TOC 2", style="TOC 2")
    document.add_paragraph("TOC 3", style="TOC 3")
    document.save(path)


class VisualValidatorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.profile = self.root / "profile.yaml"
        self.profile.write_text("visual:\n  minimum_h1_body_ratio: 1.4\n  minimum_h2_body_ratio: 1.2\n  minimum_h3_body_ratio: 1.08\n", encoding="utf-8")

    def tearDown(self):
        self.temp.cleanup()

    def report(self, broken=None):
        path = self.root / f"{broken or 'valid'}.docx"
        _make_doc(path, broken=broken)
        return visual.validate(path, self.profile)

    def test_positive_technical_whitepaper_styles_pass(self):
        result = self.report()
        self.assertNotEqual(result["status"], "INVALID")
        self.assertEqual(result["styles"]["Heading 2"]["size_pt"], 15.5)
        self.assertTrue(result["visual_review_required"])

    def test_flat_heading_sizes_are_blocked(self):
        codes = {item["code"] for item in self.report("flat")["findings"]}
        self.assertIn("HEADING_SIZE_ORDER_INVALID", codes)
        self.assertIn("HEADING_BODY_CONTRAST_TOO_LOW", codes)

    def test_inverted_heading_order_is_blocked(self):
        self.assertIn("HEADING_SIZE_ORDER_INVALID", {item["code"] for item in self.report("inverted")["findings"]})

    def test_heading_spacing_is_a_warning(self):
        self.assertIn("HEADING_SPACING_TOO_LOW", {item["code"] for item in self.report("spacing")["findings"]})

    def test_flat_toc_is_blocked(self):
        self.assertIn("TOC_HIERARCHY_INSUFFICIENT", {item["code"] for item in self.report("flat-toc")["findings"]})

    def test_ascii_font_collapse_is_blocked(self):
        self.assertIn("CODE_BODY_INDISTINGUISHABLE", {item["code"] for item in self.report("ascii-body")["findings"]})

    def test_negative_fixture_catalog_is_complete(self):
        data = json.loads((ROOT / "tests" / "fixtures" / "visual_negative_cases.json").read_text(encoding="utf-8"))
        self.assertEqual(len(data["cases"]), 8)
        self.assertIn("sparse_final_page", {item["name"] for item in data["cases"]})


if __name__ == "__main__":
    unittest.main()
