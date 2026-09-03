#!/usr/bin/env python3
"""Validate Markdown-to-DOCX semantic fidelity.

This is deliberately a publication-layer validator.  It does not rewrite the
source or repair a DOCX.  A renderer may use a different implementation, but
it must expose the same Word semantics: heading styles, explicit numbering
instances, classified code blocks, and a usable TOC cache.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from docx import Document
from lxml import etree


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
HEADING_RE = re.compile(r"^(#{1,6})[ \t]+(.+?)[ \t]*$")
LIST_RE = re.compile(r"^([ \t]*)([-+*]|\d+[.)])[ \t]+(.*)$")
FENCE_RE = re.compile(r"^([ \t]*)(`{3,}|~{3,})(.*)$")
WORD_HEADING_RE = re.compile(r"^Heading ([1-6])$")
NORMATIVE = ("MUST", "SHOULD", "MAY")


@dataclass
class Heading:
    level: int
    text: str


@dataclass
class ListItem:
    kind: str
    level: int
    marker: str
    text: str


@dataclass
class ListBlock:
    items: list[ListItem] = field(default_factory=list)

    @property
    def ordered(self) -> bool:
        return any(item.kind == "ordered" for item in self.items)

    @property
    def unordered(self) -> bool:
        return any(item.kind == "unordered" for item in self.items)


@dataclass
class SourceModel:
    headings: list[Heading] = field(default_factory=list)
    list_blocks: list[ListBlock] = field(default_factory=list)
    code_blocks: list[list[str]] = field(default_factory=list)
    ascii_diagrams: list[bool] = field(default_factory=list)
    tables: int = 0
    table_rows: list[list[list[str]]] = field(default_factory=list)
    blockquotes: int = 0
    normative: dict[str, int] = field(default_factory=dict)
    text: str = ""
    sha256: str = ""


@dataclass
class DocParagraph:
    text: str
    style: str
    level: int | None = None
    num_id: str | None = None
    ilvl: int | None = None
    number_start: int | None = None


@dataclass
class DocModel:
    paragraphs: list[DocParagraph] = field(default_factory=list)
    headings: list[Heading] = field(default_factory=list)
    list_runs: list[list[DocParagraph]] = field(default_factory=list)
    leaks: list[str] = field(default_factory=list)
    code_blocks: list[tuple[str, list[str]]] = field(default_factory=list)
    toc_fields: int = 0
    toc_entries: list[str] = field(default_factory=list)
    tables: int = 0
    table_rows: list[list[list[str]]] = field(default_factory=list)
    text: str = ""


@dataclass
class Finding:
    code: str
    message: str
    blocking: bool = True


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def canonical_text(value: str) -> str:
    value = re.sub(r"\*\*(.+?)\*\*", r"\1", value)
    value = re.sub(r"__(.+?)__", r"\1", value)
    value = re.sub(r"`([^`]+)`", r"\1", value)
    return re.sub(r"\s+", " ", value.strip())


def is_ascii_diagram(lines: list[str]) -> bool:
    if len(lines) < 2:
        return False
    value = "\n".join(lines)
    box = sum(value.count(ch) for ch in "┌┐└┘├┤┬┴┼│─╭╮╰╯")
    arrows = len(re.findall(r"(?:-->|<--|->|<-|=>|↓|↑|→|←|↔)", value))
    tree = len(re.findall(r"(?:^|\n)[ \t]*(?:[├└│]──|[+|]--|`--)", value))
    aligned = sum(1 for line in lines if line.startswith(("  ", "\t")))
    return box >= 2 or arrows >= 2 or tree >= 2 or (aligned >= 2 and len(lines) >= 3)


def parse_source(text: str) -> SourceModel:
    model = SourceModel(text=text)
    model.sha256 = hashlib.sha256(text.encode("utf-8")).hexdigest()
    lines = text.splitlines()
    outside: list[str] = []
    index = 0
    while index < len(lines):
        fence = FENCE_RE.match(lines[index])
        if fence:
            marker = fence.group(2)[0]
            code: list[str] = []
            index += 1
            while index < len(lines) and not re.match(rf"^[ \t]*{re.escape(marker)}{{3,}}", lines[index]):
                code.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            model.code_blocks.append(code)
            model.ascii_diagrams.append(is_ascii_diagram(code))
            continue
        outside.append(lines[index])
        heading = HEADING_RE.match(lines[index])
        if heading:
            model.headings.append(Heading(len(heading.group(1)), canonical_text(heading.group(2))))
        index += 1

    i = 0
    while i < len(outside):
        line = outside[i]
        item = LIST_RE.match(line)
        if item:
            block = ListBlock()
            while i < len(outside):
                current = LIST_RE.match(outside[i])
                if not current:
                    break
                indent, marker, value = current.groups()
                level = len(indent.expandtabs(4)) // 2
                kind = "ordered" if marker[0].isdigit() else "unordered"
                block.items.append(ListItem(kind, level, marker, value.strip()))
                i += 1
            model.list_blocks.append(block)
            continue
        i += 1

    for i, line in enumerate(outside):
        if line.startswith(">") and (i == 0 or not outside[i - 1].startswith(">")):
            model.blockquotes += 1
        if line.startswith("|") and i + 1 < len(outside) and re.search(r"\|?\s*:?-{3,}", outside[i + 1]):
            model.tables += 1
            rows: list[list[str]] = []
            j = i
            while j < len(outside) and outside[j].startswith("|"):
                value = outside[j].strip().strip("|")
                rows.append([canonical_text(cell.replace("\\|", "|")) for cell in value.split("|")])
                j += 1
            if len(rows) > 1:
                rows.pop(1)
            model.table_rows.append(rows)

    model.normative = {token: len(re.findall(rf"\b{token}\b", text)) for token in NORMATIVE}
    return model


def _paragraph_num_pr(paragraph) -> tuple[str | None, int | None]:
    ppr = paragraph._p.pPr
    if ppr is None or ppr.numPr is None:
        return None, None
    num_id = ppr.numPr.numId
    ilvl = ppr.numPr.ilvl
    num_value = num_id.val if num_id is not None else None
    ilvl_value = ilvl.val if ilvl is not None else None
    return (str(num_value) if num_value is not None else None, ilvl_value)


def _numbering_starts(docx_path: Path) -> dict[tuple[str, int], int]:
    starts: dict[tuple[str, int], int] = {}
    with ZipFile(docx_path) as archive:
        if "word/numbering.xml" not in archive.namelist():
            return starts
        root = etree.fromstring(archive.read("word/numbering.xml"))
        abstract: dict[tuple[str, int], int] = {}
        for node in root.xpath("./w:abstractNum", namespaces=NS):
            abstract_id = node.get(f"{{{W}}}abstractNumId")
            for level in node.xpath("./w:lvl", namespaces=NS):
                ilvl = int(level.get(f"{{{W}}}ilvl", "0"))
                start = level.xpath("./w:start/@w:val", namespaces=NS)
                abstract[(abstract_id, ilvl)] = int(start[0]) if start else 1
        for node in root.xpath("./w:num", namespaces=NS):
            num_id = node.get(f"{{{W}}}numId")
            abstract_id = node.xpath("./w:abstractNumId/@w:val", namespaces=NS)
            if not abstract_id:
                continue
            for ilvl in range(0, 9):
                value = abstract.get((abstract_id[0], ilvl), 1)
                override = node.xpath(f"./w:lvlOverride[@w:ilvl='{ilvl}']/w:startOverride/@w:val", namespaces=NS)
                starts[(num_id, ilvl)] = int(override[0]) if override else value
    return starts


def extract_docx(docx_path: Path) -> DocModel:
    document = Document(str(docx_path))
    starts = _numbering_starts(docx_path)
    model = DocModel(tables=len(document.tables))
    model.table_rows = [[[cell.text.strip() for cell in row.cells] for row in table.rows] for table in document.tables]
    previous_was_list = False
    current_run: list[DocParagraph] = []
    for paragraph in document.paragraphs:
        style = paragraph.style.name if paragraph.style else ""
        num_id, ilvl = _paragraph_num_pr(paragraph)
        heading = WORD_HEADING_RE.match(style)
        level = int(heading.group(1)) if heading else None
        is_list = style.lower().startswith("list ") or num_id is not None
        kind = "ordered" if ("number" in style.lower() or (num_id is not None and not "bullet" in style.lower())) else "unordered"
        if is_list:
            record = DocParagraph(
                paragraph.text,
                style,
                level=None,
                num_id=num_id or f"style:{style}",
                ilvl=ilvl if ilvl is not None else 0,
                number_start=starts.get((num_id, ilvl or 0)) if num_id else None,
            )
            record.style = f"{kind}:{style}"
            if not previous_was_list and current_run:
                model.list_runs.append(current_run)
                current_run = []
            current_run.append(record)
            previous_was_list = True
        else:
            if current_run:
                model.list_runs.append(current_run)
                current_run = []
            previous_was_list = False

        record = DocParagraph(paragraph.text, style, level, num_id, ilvl)
        model.paragraphs.append(record)
        if heading:
            model.headings.append(Heading(level, canonical_text(paragraph.text)))
        if re.match(r"^#{1,6}[ \t]+", paragraph.text) and not style.startswith(("Code", "ASCII Diagram", "Literal")):
            model.leaks.append(paragraph.text)
        if style in ("Code", "ASCII Diagram"):
            model.code_blocks.append((style, paragraph.text.split("\n")))
        model.text += paragraph.text + "\n"
    if current_run:
        model.list_runs.append(current_run)

    toc_seen = False
    for index, paragraph in enumerate(document.paragraphs):
        xml = paragraph._p.xml
        if re.search(r"\bTOC\b", xml):
            model.toc_fields += 1
            toc_seen = True
            continue
        if toc_seen:
            if 'w:type="page"' in xml:
                break
            value = paragraph.text.strip()
            if value and value != "目录":
                model.toc_entries.append(value)
    return model


def _source_summary(source: SourceModel) -> dict[str, Any]:
    return {
        "H1": sum(h.level == 1 for h in source.headings),
        "H2": sum(h.level == 2 for h in source.headings),
        "H3": sum(h.level == 3 for h in source.headings),
        "H4": sum(h.level == 4 for h in source.headings),
        "H5": sum(h.level == 5 for h in source.headings),
        "H6": sum(h.level == 6 for h in source.headings),
        "ordered_list_blocks": sum(block.ordered for block in source.list_blocks),
        "unordered_list_blocks": sum(block.unordered for block in source.list_blocks),
        "fenced_code_blocks": len(source.code_blocks),
        "ascii_diagram_blocks": sum(source.ascii_diagrams),
        "tables": source.tables,
        "blockquotes": source.blockquotes,
    }


def _doc_summary(doc: DocModel) -> dict[str, Any]:
    return {
        **{f"H{i}": sum(h.level == i for h in doc.headings) for i in range(1, 7)},
        "list_number_paragraphs": sum(p.style.startswith("ordered:") for run in doc.list_runs for p in run),
        "list_bullet_paragraphs": sum(p.style.startswith("unordered:") for run in doc.list_runs for p in run),
        "code_style_blocks": sum(style == "Code" for style, _ in doc.code_blocks),
        "ascii_diagram_style_blocks": sum(style == "ASCII Diagram" for style, _ in doc.code_blocks),
        "literal_markdown_heading_leaks": len(doc.leaks),
        "toc_fields": doc.toc_fields,
        "toc_entries": len(doc.toc_entries),
        "tables": doc.tables,
    }


def _profile_value(profile: Path | None, key: str, default: Any, section: str | None = None) -> Any:
    if not profile or not profile.exists():
        return default
    text = profile.read_text(encoding="utf-8")
    if section:
        section_match = re.search(rf"^{re.escape(section)}:\s*$", text, re.MULTILINE)
        if section_match:
            tail = text[section_match.end():]
            next_section = re.search(r"^\S.*:\s*$", tail, re.MULTILINE)
            text = tail[: next_section.start()] if next_section else tail
    match = re.search(rf"^\s+{re.escape(key)}:\s*[\"']?([^\"'\n#]+?)[\"']?\s*$", text, re.MULTILINE)
    if not match:
        return default
    value = match.group(1).strip()
    if value.lower() in ("true", "false"):
        return value.lower() == "true"
    try:
        return int(value)
    except ValueError:
        return value


def validate(source_path: Path, docx_path: Path, profile_path: Path | None = None, expected_sha: str | None = None, strict: bool = False) -> dict[str, Any]:
    source = parse_source(source_path.read_text(encoding="utf-8"))
    document = extract_docx(docx_path)
    findings: list[Finding] = []

    if expected_sha and source.sha256 != expected_sha:
        findings.append(Finding("SOURCE_CHANGED_DURING_PUBLICATION", f"expected {expected_sha}, got {source.sha256}"))

    if len(document.headings) != len(source.headings):
        findings.append(Finding("HEADING_LEVEL_MISMATCH", f"heading count expected {len(source.headings)}, got {len(document.headings)}"))
    for index, expected in enumerate(source.headings[: len(document.headings)]):
        actual = document.headings[index]
        if expected.level != actual.level:
            findings.append(Finding("HEADING_LEVEL_MISMATCH", f"heading {index + 1}: expected H{expected.level}, got H{actual.level}"))
        if expected.text != actual.text:
            findings.append(Finding("HEADING_TEXT_MISSING", f"heading {index + 1}: expected {expected.text!r}, got {actual.text!r}"))
    source_jumps = {
        index
        for index, (previous, current) in enumerate(zip(source.headings, source.headings[1:]), start=1)
        if current.level > previous.level + 1
    }
    if source_jumps:
        findings.append(Finding("SOURCE_HEADING_HIERARCHY_WARNING", f"source contains {len(source_jumps)} heading level jump(s); preserved without source mutation", blocking=False))
    for index, (previous, current) in enumerate(zip(document.headings, document.headings[1:]), start=1):
        if current.level > previous.level + 1 and index not in source_jumps:
            findings.append(Finding("HEADING_HIERARCHY_INVALID", f"DOCX introduced illegal jump H{previous.level} -> H{current.level}"))
            break

    if document.leaks:
        findings.append(Finding("LITERAL_MARKDOWN_HEADING_LEAK", f"{len(document.leaks)} body paragraphs contain Markdown heading markers"))

    source_lists = source.list_blocks
    doc_lists = document.list_runs
    if len(source_lists) != len(doc_lists):
        findings.append(Finding("LIST_RESTART_MISMATCH", f"list block count expected {len(source_lists)}, got {len(doc_lists)}"))
    for index, expected in enumerate(source_lists[: len(doc_lists)]):
        actual = doc_lists[index]
        expected_kinds = [item.kind for item in expected.items]
        actual_kinds = ["ordered" if item.style.startswith("ordered:") else "unordered" for item in actual]
        expected_levels = [item.level for item in expected.items]
        actual_levels = [item.ilvl or 0 for item in actual]
        expected_texts = [canonical_text(item.text) for item in expected.items]
        actual_texts = [canonical_text(item.text) for item in actual]
        if len(expected_kinds) != len(actual_kinds) or expected_kinds != actual_kinds:
            findings.append(Finding("LIST_RESTART_MISMATCH", f"list block {index + 1}: item/type structure differs"))
        if expected_levels != actual_levels:
            findings.append(Finding("LIST_LEVEL_MISMATCH", f"list block {index + 1}: expected levels {expected_levels}, got {actual_levels}"))
        if expected_texts != actual_texts:
            findings.append(Finding("LIST_ITEM_TEXT_MISSING", f"list block {index + 1}: item text differs"))
        if expected.ordered:
            first = next(item for item in expected.items if item.kind == "ordered")
            actual_first = next((item for item in actual if item.style.startswith("ordered:")), None)
            if actual_first and actual_first.number_start is not None and int(re.match(r"\d+", first.marker).group()) != actual_first.number_start:
                findings.append(Finding("LIST_RESTART_MISMATCH", f"list block {index + 1}: expected start {first.marker}, got {actual_first.number_start}"))
    ordered_ids: list[str] = []
    for run in doc_lists:
        if any(item.style.startswith("ordered:") for item in run):
            ordered_ids.append(next(item.num_id or "" for item in run if item.style.startswith("ordered:")))
    for left, right in zip(ordered_ids, ordered_ids[1:]):
        if left == right:
            findings.append(Finding("UNEXPECTED_LIST_CONTINUATION", f"independent ordered lists share numbering instance {left}"))
            break

    expected_code_count = len(source.code_blocks)
    if len(document.code_blocks) != expected_code_count:
        findings.append(Finding("CODE_INDENTATION_CORRUPTED", f"code/diagram block count expected {expected_code_count}, got {len(document.code_blocks)}"))
    for index, (expected_lines, expected_ascii) in enumerate(zip(source.code_blocks, source.ascii_diagrams)):
        if index >= len(document.code_blocks):
            break
        style, actual_lines = document.code_blocks[index]
        expected_style = "ASCII Diagram" if expected_ascii else "Code"
        if style != expected_style:
            findings.append(Finding("ASCII_DIAGRAM_CLASS_MISMATCH", f"block {index + 1}: expected {expected_style}, got {style}"))
        if expected_lines != actual_lines:
            findings.append(Finding("ASCII_DIAGRAM_WRAPPED" if expected_ascii else "CODE_INDENTATION_CORRUPTED", f"block {index + 1}: line breaks or whitespace changed"))

    for token, expected_count in source.normative.items():
        actual_count = len(re.findall(rf"\b{token}\b", document.text))
        if actual_count != expected_count:
            findings.append(Finding("CONTENT_FIDELITY_MISMATCH", f"{token} count expected {expected_count}, got {actual_count}"))
    if document.tables != source.tables:
        findings.append(Finding("TABLE_ROW_LOSS", f"table count expected {source.tables}, got {document.tables}"))
    if source.table_rows != document.table_rows:
        findings.append(Finding("TABLE_ROW_LOSS", "table cell or row content differs from Markdown"))

    toc_enabled = _profile_value(profile_path, "enabled", True, "toc")
    toc_depth = _profile_value(profile_path, "depth", _profile_value(profile_path, "toc_depth", 3, "toc"), "toc")
    if toc_enabled:
        if document.toc_fields == 0:
            findings.append(Finding("TOC_EMPTY", "TOC field is missing"))
        expected_toc = [h.text for h in source.headings if h.level <= int(toc_depth)]
        missing = [title for title in expected_toc if not any(title in entry for entry in document.toc_entries)]
        if not document.toc_entries:
            findings.append(Finding("TOC_EMPTY", "TOC has no visible cached entries"))
        elif missing:
            findings.append(Finding("TOC_INCOMPLETE", f"TOC is missing {len(missing)} expected entries"))
        if len(document.toc_entries) <= 1 and len(expected_toc) > 1:
            findings.append(Finding("TOC_INCOMPLETE", "TOC contains only a title or one cached entry"))

    blocking = [finding for finding in findings if finding.blocking]
    warnings = [finding for finding in findings if not finding.blocking]
    status = "INVALID" if blocking else ("VALID_WITH_WARNINGS" if warnings or strict and findings else "VALID")
    return {
        "status": status,
        "source": _source_summary(source),
        "docx": _doc_summary(document),
        "source_sha256": source.sha256,
        "findings": [asdict(finding) for finding in findings],
    }


def print_report(report: dict[str, Any], as_json: bool = False) -> None:
    if as_json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return
    print("Publication Structural Fidelity Report")
    print("\nSource:")
    for key, value in report["source"].items():
        print(f"  {key}: {value}")
    print("\nDOCX:")
    for key, value in report["docx"].items():
        print(f"  {key}: {value}")
    print("\nFindings:")
    if not report["findings"]:
        print("  none")
    else:
        for finding in report["findings"]:
            marker = "BLOCKING" if finding["blocking"] else "WARNING"
            print(f"  [{marker}] {finding['code']}: {finding['message']}")
    print(f"\nStatus: {report['status']}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--docx", required=True, type=Path)
    parser.add_argument("--profile", type=Path)
    parser.add_argument("--source-sha256")
    parser.add_argument("--strict", action="store_true")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)
    report = validate(args.source, args.docx, args.profile, args.source_sha256, args.strict)
    print_report(report, args.json)
    return 1 if report["status"] == "INVALID" else 0


if __name__ == "__main__":
    raise SystemExit(main())
